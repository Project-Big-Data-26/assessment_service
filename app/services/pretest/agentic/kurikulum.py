"""
kurikulum.py -- acuan TUNGGAL struktur bab, subbab, dan prasyarat dari
data/bab_prasyarat/Prasyarat_Math.xlsx (dibuat tim; jangan ditebak dari buku).

  sheet "Sheet2"  : daftar subbab (Jenjang, Kelas, Jalur SMA, No. Bab, Nama Bab, Kode, Nama Subbab)
  sheet "Per BAB" : bab_id ("SD5-B7", "SMA11L-B2") + prasyarat_bab_id ("SD4-B5, SD5-B6")

Pemetaan kunci buku pipeline -> awalan bab_id:
  kelas1..6 -> SD1..SD6, kelas7..9 -> SMP7..SMP9, kelas10 -> SMA10,
  kelas11 -> SMA11U, kelas11_tl -> SMA11L, kelas12 -> SMA12U, kelas12_tl -> SMA12L
"""

from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from functools import lru_cache
from typing import Dict, List, Optional, Tuple

from app.services.pretest.agentic.books import DATA_DIR

EXCEL_PATH = os.path.join(DATA_DIR, "bab_prasyarat", "Prasyarat_Math.xlsx")


@dataclass
class Bab:
    bab_id: str                               # "SD5-B7"
    book: str                                 # "kelas5"
    bab_no: int
    nama: str
    subbab: List[Tuple[str, str]] = field(default_factory=list)   # [("A", "Membandingkan Ciri-Ciri Segitiga"), ...]
    prasyarat: List[str] = field(default_factory=list)            # ["SD4-B5", "SD5-B6"]

    def subbab_nama(self, kode: str) -> Optional[str]:
        return next((n for k, n in self.subbab if k == kode), None)

    @property
    def subbab_terakhir(self) -> Optional[Tuple[str, str]]:
        return self.subbab[-1] if self.subbab else None


def book_prefix(book: str) -> str:
    kelas = int(re.search(r"\d+", book).group())
    lanjut = book.endswith("_tl")
    if kelas <= 6:
        return f"SD{kelas}"
    if kelas <= 9:
        return f"SMP{kelas}"
    if kelas == 10:
        return "SMA10"
    return f"SMA{kelas}{'L' if lanjut else 'U'}"


def book_of_prefix(prefix: str) -> str:
    m = re.match(r"^(SD|SMP|SMA)(\d+)([UL]?)$", prefix)
    kelas, jalur = int(m.group(2)), m.group(3)
    return f"kelas{kelas}" + ("_tl" if jalur == "L" else "")


def bab_id(book: str, bab_no: int) -> str:
    return f"{book_prefix(book)}-B{bab_no}"


def _clean(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip()


@lru_cache(maxsize=1)
def load(path: str = EXCEL_PATH) -> Dict[str, Bab]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    babs: Dict[str, Bab] = {}
    for r in wb["Sheet2"].iter_rows(min_row=2, values_only=True):
        jenjang, kelas, jalur, no, nama, kode, sub = r[1:8]   # kolom 0 = "No"
        if not jenjang or kelas is None or no is None or not kode:
            continue
        kelas, no = int(kelas), int(no)
        suffix = ""
        if jenjang == "SMA" and kelas >= 11:
            suffix = "L" if "lanjut" in str(jalur).lower() else "U"
        prefix = f"{jenjang}{kelas}{suffix}"
        bid = f"{prefix}-B{no}"
        bab = babs.setdefault(bid, Bab(bid, book_of_prefix(prefix), no, _clean(nama)))
        bab.subbab.append((_clean(kode).upper(), _clean(sub)))
    for r in wb["Per BAB"].iter_rows(values_only=True):
        if len(r) < 6:   # baris kosong di mode read_only bisa lebih pendek
            continue
        bid, pras = _clean(r[1]), _clean(r[5])
        if bid in babs and pras and pras != "-":
            babs[bid].prasyarat = [p.strip() for p in pras.split(",") if p.strip()]
    wb.close()
    return babs


def get(book: str, bab_no: int) -> Optional[Bab]:
    return load().get(bab_id(book, bab_no))
